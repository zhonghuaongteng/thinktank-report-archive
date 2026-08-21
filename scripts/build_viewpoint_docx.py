from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "687386"
LIGHT = "F2F4F7"
CALLOUT = "F4F6F9"


def set_font(run, name="Microsoft YaHei", size=11, bold=False, color="202124"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_hyperlink(paragraph, label, url):
    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([rpr, text])
    link.append(run)
    paragraph._p.append(link)


def add_markdown_runs(paragraph, text):
    pos = 0
    pattern = re.compile(r"\[([^]]+)\]\((https?://[^)]+)\)|\*\*([^*]+)\*\*")
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_font(paragraph.add_run(text[pos:match.start()]))
        if match.group(1):
            add_hyperlink(paragraph, match.group(1), match.group(2))
        else:
            set_font(paragraph.add_run(match.group(3)), bold=True, color=INK)
        pos = match.end()
    if pos < len(text):
        set_font(paragraph.add_run(text[pos:]))


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    set_font(run, size=9, color=MUTED)


def setup_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(hp.add_run("国际科技智库观点演变研究 | 2016—2026"), size=8.5, color=MUTED)
    ppr = hp._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), "D9DEE7")
    pbdr.append(bottom)
    ppr.append(pbdr)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    set_font(p.add_run("研究专报"), size=11, bold=True, color="7A5A00")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run("从开放创新到受控互赖"), size=28, bold=True, color=INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(32)
    set_font(p.add_run("国际主要科技智库十年科技战略议程的目标重排与工具硬化"), size=14, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("材料范围：2016年1月1日至2026年8月21日"), size=10.5, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("研究时点：2026年8月21日"), size=10.5, color=MUTED)
    doc.add_page_break()


def add_callout(doc, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "E1E6EE")
        edge.set(qn("w:space"), "8")
        p_bdr.append(edge)
    p_pr.append(p_bdr)
    set_font(p.add_run(title), size=11, bold=True, color=INK)
    p.add_run().add_break()
    add_markdown_runs(p, body)


def add_table(doc, rows):
    cols = len(rows[0])
    widths = [1500, 2750, 2900, 2210] if cols == 4 else [9360 // cols] * cols
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, row in enumerate(rows):
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(value), size=8.5 if i else 9, bold=(i == 0), color=INK if i == 0 else "202124")
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def render_markdown(doc, lines):
    i = 0
    first_h1 = True
    while i < len(lines):
        raw = lines[i].rstrip()
        if not raw or raw.startswith("**研究时点") or raw.startswith("**材料范围"):
            i += 1
            continue
        if raw.startswith("# "):
            if first_h1:
                first_h1 = False
                i += 1
                continue
            doc.add_heading(raw[2:], level=1)
        elif raw.startswith("## "):
            doc.add_heading(raw[3:], level=1)
        elif raw.startswith("### "):
            doc.add_heading(raw[4:], level=2)
        elif raw.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                parts = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(set(x) <= {"-", ":"} for x in parts):
                    rows.append(parts)
                i += 1
            add_table(doc, rows)
            continue
        elif re.match(r"^\d+\. ", raw):
            p = doc.add_paragraph(style="List Number")
            add_markdown_runs(p, re.sub(r"^\d+\. ", "", raw))
        elif raw.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_markdown_runs(p, raw[2:])
        else:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Inches(0.28)
            add_markdown_runs(p, raw)
        i += 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--output", required=True, type=Path)
    args = ap.parse_args()
    lines = args.input.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_styles(doc)
    setup_section(doc.sections[0])
    add_cover(doc)
    add_callout(
        doc,
        "核心判断",
        "十年变化表现为科技政策目标从增长与创新效率扩展到转型、韧性与安全，政策工具由通用投入转向使命、关键产业、供应链和研究安全；国际合作被改写为可选择、可审查、可治理的互赖。",
    )
    render_markdown(doc, lines)
    props = doc.core_properties
    props.title = "从开放创新到受控互赖"
    props.subject = "国际主要科技智库观点演变研究（2016—2026）"
    props.author = "研究知识库"
    props.keywords = "科技智库, 科技战略, 受控互赖, 产业政策, 研究安全"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
